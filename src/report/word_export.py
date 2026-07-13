"""Export business research report to Microsoft Word (.docx)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from src.config import PROJECT_ROOT, get_path, load_config
from src.dashboard.data_loader import load_live_forecast

SECTOR_EN = {
    "通信": "Telecommunications",
    "有色金属": "Non-ferrous Metals",
    "传媒": "Media",
    "国防军工": "Defence",
    "机械设备": "Machinery",
    "家用电器": "Home Appliances",
    "建筑装饰": "Construction & Decoration",
    "交通运输": "Transportation",
    "农林牧渔": "Agriculture",
    "汽车": "Automobiles",
}


def _pct(v, d: int = 1) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "—"
    return f"{float(v) * 100:.{d}f}%"


def _num(v, d: int = 2) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "—"
    return f"{float(v):.{d}f}"


def _sector_label(name: str) -> str:
    en = SECTOR_EN.get(name)
    return f"{name} ({en})" if en else name


def _read_csv(rel: str) -> pd.DataFrame | None:
    path = get_path("processed") / rel
    return pd.read_csv(path) if path.exists() else None


def _set_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def _add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(30, 58, 95)


def _add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(90, 100, 115)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_body(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def _add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            table.rows[r_idx].cells[c_idx].text = str(val)
    doc.add_paragraph()


def _add_figure(doc: Document, rel_path: str, caption: str, width: float = 6.0) -> None:
    path = (PROJECT_ROOT / rel_path).resolve()
    if not path.exists():
        _add_body(doc, f"[Figure unavailable: {caption}]")
        return
    doc.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 100, 115)


def _business_conclusion_paragraphs(fc: dict) -> list[str]:
    holdings = fc.get("top_holdings", [])
    names = ", ".join(_sector_label(h["industry"]) for h in holdings[:5])
    forecast_month = fc.get("forecast_month", "the target month")
    data_as_of = fc.get("data_as_of", "the latest feature month")

    hits = sum(
        1 for h in holdings
        if h.get("actual_1m") is not None and h["actual_1m"] > 0
    )
    n_hold = len(holdings)
    top = holdings[0] if holdings else {}
    top_name = _sector_label(top.get("industry", ""))
    top_actual = top.get("actual_1m")

    paras = [
        (
            f"Applied conclusion ({forecast_month}): Using information available through "
            f"{data_as_of}, the production model assigns the highest relative-return potential "
            f"to {names}. This is not presented as a forecast of certain outperformance, but as "
            f"a ranked starting point for monthly sector allocation — the same output an asset "
            f"allocation team would require before a committee discussion."
        ),
        (
            "The appropriate business reading is probabilistic. Over the full out-of-sample "
            "window, the Top-5 hit rate is 25.5%: roughly one in four recommended sectors "
            "finishes in the realised Top-5 next month. That is modest in absolute terms, yet "
            "materially above a random null (near 18.5%) and sufficient to justify use as a "
            "structured decision input rather than a discretionary shortcut."
        ),
    ]

    if top_actual is not None and n_hold:
        paras.append(
            f"In {forecast_month}, the model's top-ranked sector ({top_name}) returned "
            f"{_pct(top_actual)} ex post, while {hits} of {n_hold} overweight names delivered "
            f"positive monthly returns. A single month neither validates nor refutes the "
            f"framework; it illustrates why the system is designed for repeat application, "
            f"regime monitoring, and benchmark cross-check — not for one-shot prediction claims."
        )

    paras.append(
        "For implementation, the recommended overweight set should be combined with two "
        "governance checks already embedded in the project: (i) overlap with the 12-month "
        "momentum Top-5, where sectors appearing in both lists receive higher conviction; "
        "and (ii) current market regime diagnostics, reducing active risk when historical "
        "failure rates were elevated (e.g. weak momentum combined with low volatility in 2023)."
    )
    return paras


def _why_project_matters() -> list[str]:
    return [
        (
            "Reproducibility at zero data cost. The entire pipeline runs on public AkShare "
            "inputs. Any stakeholder can audit how a recommendation was produced, which "
            "variables entered the model, and which publication lags were applied — a property "
            "commercial databases do not automatically confer."
        ),
        (
            "Leakage-controlled validation. Nested walk-forward evaluation selects "
            "hyperparameters on a rolling 12-month validation window before each out-of-sample "
            "month. This design choice prioritises credibility over in-sample fit, which is "
            "the difference between a research artefact and a decision-support tool."
        ),
        (
            "Explainability for committee settings. SHAP attribution and sector-level case "
            "studies allow the research function to answer why a sector ranks first — a "
            "requirement in institutional workflows that pure momentum rules cannot satisfy "
            "without additional narrative work."
        ),
        (
            "Honest benchmark disclosure. The strategy outperforms the CSI 300 (Sharpe 0.47 "
            "vs 0.25) and sector equal weight, but trails a simple 12-month momentum Top-5 "
            "(Sharpe 0.53). Publishing this gap is a feature, not a flaw: it defines where "
            "the model adds value (process, integration, governance) and where simpler rules "
            "remain competitive."
        ),
        (
            "Productised delivery. The Streamlit decision dashboard converts model output into "
            "executive KPIs, allocation weights, regime context, and drill-down sector views — "
            "bridging quantitative research and business analytics without requiring Python "
            "literacy from end users."
        ),
    ]


def generate_word(
    output_path: Path | None = None,
    include_figures: bool = True,
) -> Path:
    output_path = output_path or PROJECT_ROOT / "docs" / "REPORT.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    fc = load_live_forecast()
    primary = _read_csv("ml/model_comparison_primary.csv")
    bench = _read_csv("ml/benchmark_comparison.csv")
    ic_decay = _read_csv("ml/analysis/ic_decay_summary.csv")
    subperiod = _read_csv("ml/subperiod_analysis.csv")
    shap = _read_csv("ml/feature_selection/shap_importance.csv")

    metrics = primary.iloc[0].to_dict() if primary is not None and not primary.empty else {}
    report_date = datetime.now().strftime("%d %B %Y")
    project_name = load_config().get("project", {}).get(
        "name", "Public-Data Sector Rotation Decision Support System"
    )

    doc = Document()
    _set_style(doc)

    # ── Cover ──────────────────────────────────────────────────────────────
    _add_title(doc, project_name)
    _add_subtitle(doc, "Decision Support System · Research Report")
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Report Date: {report_date}\n").font.size = Pt(10)
    meta.add_run("Sample: January 2020 – May 2026  |  OOS: 40 months  |  Public data only").font.size = Pt(10)
    doc.add_page_break()

    # ── Executive Summary ───────────────────────────────────────────────────
    _add_heading(doc, "Executive Summary", 1)
    _add_body(
        doc,
        "This study asks whether publicly available information can support monthly "
        "cross-sectional ranking of Shenwan Level-1 industries — not point forecasts of "
        "absolute return, but a repeatable answer to which sectors are likely to outperform "
        "on a relative basis next month.",
    )
    _add_body(
        doc,
        f"Out-of-sample over 40 months, the LambdaRank + Softmax Top-5 strategy delivers "
        f"a net annualised return of {_pct(metrics.get('annual_return'))}, Sharpe {_num(metrics.get('sharpe'))}, "
        f"and mean Rank IC {_num(metrics.get('rank_ic_mean'), 3)}. The model beats passive "
        f"benchmarks but does not surpass a 12-month momentum rule. The evidence supports "
        f"deployment as decision support — not as a standalone alpha engine.",
    )

    # ── Business Conclusion (core ask) ──────────────────────────────────────
    _add_heading(doc, "Business Conclusion — How the Model Is Applied", 1)
    if fc:
        for para in _business_conclusion_paragraphs(fc):
            _add_body(doc, para)

        _add_heading(doc, "Recommended Overweight Set", 2)
        rows = []
        for h in fc.get("top_holdings", []):
            actual = _pct(h["actual_1m"]) if h.get("actual_1m") is not None else "Pending"
            rows.append([
                str(h["rank"]),
                _sector_label(h["industry"]),
                _num(h["score"], 3),
                _pct(h["weight"]),
                actual,
            ])
        _add_table(
            doc,
            ["Rank", "Sector (Shenwan L1)", "Score", "Weight", "Realised 1M Return"],
            rows,
        )

        under = [r for r in fc.get("rankings", []) if r.get("signal") == "Underweight"]
        if under:
            under_names = ", ".join(_sector_label(r["industry"]) for r in under[:5])
            _add_body(
                doc,
                f"Underweight candidates at the bottom of the cross-section: {under_names}. "
                f"These names are deprioritised relative to the broader universe, not necessarily "
                f"short candidates.",
            )
    else:
        _add_body(doc, "No live forecast available. Run run_phase1.py and run_ml.py to populate.")

    # ── Why the project matters ─────────────────────────────────────────────
    _add_heading(doc, "Why This Project Is Valuable — Despite Imperfect Accuracy", 1)
    _add_body(
        doc,
        "A decision-support system is judged not only by whether it wins every month, but by "
        "whether it improves the quality, transparency, and auditability of recurring "
        "allocation decisions. On that criterion, the project delivers five concrete benefits:",
    )
    for item in _why_project_matters():
        _add_bullet(doc, item)

    # ── Methodology (concise) ───────────────────────────────────────────────
    _add_heading(doc, "Research Design", 1)
    _add_body(
        doc,
        "Data are sourced from AkShare (macro, liquidity, market state, Shenwan L1 returns). "
        "Features at month t predict sector returns at month t+1, with publication lags enforced. "
        "Each out-of-sample month uses nested walk-forward validation: inner train, 12-month "
        "validation for hyperparameter selection, retrain on full history, then test. Portfolio "
        "construction applies Softmax weighting to the Top-5 sectors with a 20% single-name cap "
        "and 23 bps one-way transaction costs.",
    )

    # ── Performance ─────────────────────────────────────────────────────────
    _add_heading(doc, "Out-of-Sample Evidence", 1)
    _add_heading(doc, "Primary Strategy Metrics", 2)
    _add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Net Annual Return", _pct(metrics.get("annual_return"))],
            ["Sharpe Ratio", _num(metrics.get("sharpe"))],
            ["Maximum Drawdown", _pct(metrics.get("max_drawdown"))],
            ["Mean Rank IC", _num(metrics.get("rank_ic_mean"), 3)],
            ["ICIR", _num(metrics.get("icir"), 3)],
            ["Top-5 Hit Rate", _pct(metrics.get("top_k_hit_rate"))],
            ["Avg Monthly Turnover", _pct(metrics.get("avg_turnover"))],
        ],
    )

    if bench is not None and not bench.empty:
        _add_heading(doc, "Benchmark Comparison", 2)
        brow = []
        for _, r in bench.iterrows():
            brow.append([
                str(r.get("benchmark", "")),
                _pct(r.get("annual_return")),
                _num(r.get("sharpe")),
                _pct(r.get("max_drawdown")),
            ])
        _add_table(doc, ["Benchmark", "Annual Return", "Sharpe", "Max Drawdown"], brow)
        _add_body(
            doc,
            "The strategy exceeds the CSI 300, sector equal weight, and random Top-K null. "
            "It underperforms the 12-month momentum Top-5 on Sharpe. Any business deployment "
            "should treat momentum overlap and regime filters as first-class governance rules.",
        )

    if subperiod is not None and not subperiod.empty:
        _add_heading(doc, "Sub-Period Sharpe", 2)
        srows = [[str(int(r["period"])), str(int(r["n_months"])), _num(r.get("sharpe"))] for _, r in subperiod.iterrows()]
        _add_table(doc, ["Period", "Months", "Sharpe"], srows)
        _add_body(
            doc,
            "Performance was negative in 2023 (Sharpe −1.61), recovered in 2024, and strengthened "
            "in 2025. Regime sensitivity is real; the model should be scaled down in historically "
            "adverse environments rather than applied mechanically.",
        )

    if ic_decay is not None and not ic_decay.empty:
        _add_heading(doc, "Signal Persistence (IC Decay)", 2)
        irows = []
        for _, r in ic_decay.sort_values("horizon_months").iterrows():
            irows.append([
                f"{int(r['horizon_months'])}M",
                _num(r.get("rank_ic_mean"), 3),
                _num(r.get("icir"), 3),
            ])
        _add_table(doc, ["Horizon", "Mean Rank IC", "ICIR"], irows)
        _add_body(
            doc,
            "Ranking signal concentrates at the one-month horizon and decays thereafter, "
            "supporting monthly — not quarterly — rebalancing.",
        )

    if include_figures:
        _add_heading(doc, "Charts", 2)
        _add_figure(doc, "data/figures/ml/ml_nav_curves.png", "Figure 1. Cumulative NAV — Out-of-Sample")
        _add_figure(doc, "data/figures/ml/ml_ic_decay.png", "Figure 2. IC Decay by Horizon")
        _add_figure(doc, "data/figures/ml/ml_feature_importance.png", "Figure 3. Feature Importance (SHAP)")

    # ── Interpretation ────────────────────────────────────────────────────
    _add_heading(doc, "Model Interpretation", 1)
    _add_body(
        doc,
        "SHAP analysis shows 12-month relative momentum, relative strength, and price momentum "
        "as the dominant contributors. Macro and liquidity variables supplement but do not "
        "dominate. The model therefore integrates price information with public macro context — "
        "it does not claim to discover a non-momentum edge, and reporting should not imply otherwise.",
    )
    if shap is not None and not shap.empty:
        srows = []
        for i, (_, r) in enumerate(shap.head(8).iterrows(), 1):
            srows.append([str(i), str(r["feature"]), _num(r["mean_abs_shap"], 4)])
        _add_table(doc, ["Rank", "Feature", "Mean |SHAP|"], srows)

    # ── Limitations ─────────────────────────────────────────────────────────
    _add_heading(doc, "Limitations and Appropriate Use", 1)
    limits = [
        "Forty months of OOS data limit statistical confidence; sub-period results are volatile.",
        "Public-data coverage is thinner than institutional databases; fundamentals are proxy-based.",
        "Top-5 hit rate near 25% means most individual calls will miss — process value exceeds single-month accuracy.",
        "The model trails simple momentum on Sharpe; blended governance is recommended.",
        "This document is for research and business analytics demonstration only; not investment advice.",
    ]
    for lim in limits:
        _add_bullet(doc, lim)

    # ── Reproducibility ─────────────────────────────────────────────────────
    _add_heading(doc, "Reproducibility", 1)
    _add_body(doc, "python run_phase1.py --skip-fetch")
    _add_body(doc, "python run_ml.py")
    _add_body(doc, "python generate_word.py")

    doc.save(str(output_path))
    return output_path


if __name__ == "__main__":
    out = generate_word()
    print(f"Word report saved: {out}")
